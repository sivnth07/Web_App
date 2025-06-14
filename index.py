import pdfplumber
import docx
import re
import openpyxl
import os

# Function to extract email
def extract_email(text):
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email = re.findall(email_pattern, text)
    return email[0] if email else None

# Function to extract phone number
def extract_phone(text):
    phone_pattern = r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
    phone = re.findall(phone_pattern, text)
    return phone[0] if phone else None

# Function to extract total experience
def extract_experience(text):
    exp_pattern = r'(\d{1,2})\+?\s*years of experience'
    experience = re.findall(exp_pattern, text.lower())
    return experience[0] if experience else None

# Function to extract universities
def extract_universities(text):
    universities = []
    for line in text.split('\n'):
        if "university" in line.lower() or "institute" in line.lower():
            universities.append(line.strip())
    return ', '.join(universities)

# Function to extract degrees
def extract_degrees(text):
    degree_keywords = ["Bachelor", "Master", "PhD", "BSc", "MSc", "MBA"]
    degrees = [degree for degree in degree_keywords if degree.lower() in text.lower()]
    return ', '.join(degrees)

# Function to extract designation (job titles)
def extract_designation(text):
    designation_keywords = ["Engineer", "Manager", "Developer", "Consultant"]
    designations = [title for title in designation_keywords if title.lower() in text.lower()]
    return ', '.join(designations)

# Function to extract skills
def extract_skills(text):
    skill_keywords = ["Python", "Java", "SQL", "Data Analysis", "Machine Learning", "AWS"]
    skills = [skill for skill in skill_keywords if skill.lower() in text.lower()]
    return ', '.join(skills)

# Function to extract companies worked for
def extract_companies(text):
    company_keywords = ["Inc", "Ltd", "Corporation", "Company", "LLC"]
    companies = []
    for line in text.split('\n'):
        if any(keyword in line for keyword in company_keywords):
            companies.append(line.strip())
    return ', '.join(companies)

# Function to read PDF resumes
def read_pdf(file_path):
    with pdfplumber.open(file_path) as pdf:
        text = ''
        for page in pdf.pages:
            text += page.extract_text()
    return text

# Function to read DOCX resumes
def read_docx(file_path):
    doc = docx.Document(file_path)
    text = '\n'.join([para.text for para in doc.paragraphs])
    return text

# Function to process resumes and extract data
def process_resumes(folder_path):
    resume_data = []
    for file_name in os.listdir(folder_path):
        if file_name.endswith('.pdf'):
            text = read_pdf(os.path.join(folder_path, file_name))
        elif file_name.endswith('.docx'):
            text = read_docx(os.path.join(folder_path, file_name))
        else:
            print(f"Unsupported file format: {file_name}")
            continue

        # Extract information
        email = extract_email(text)
        phone = extract_phone(text)
        experience = extract_experience(text)
        universities = extract_universities(text)
        degrees = extract_degrees(text)
        designation = extract_designation(text)
        skills = extract_skills(text)
        companies = extract_companies(text)

        # Append the extracted data to a list
        resume_data.append([email, phone, experience, universities, degrees, designation, skills, companies, text])

    return resume_data

# Function to create and save data into an Excel file
def create_excel(data, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resume Data"

    headers = ["Email", "Phone", "Experience", "Universities", "Degrees", "Designation", "Skills", "Companies", "Summary"]
    ws.append(headers)

    for row in data:
        ws.append(row)

    wb.save(output_file)

# Main function to tie everything together
def main():
    folder_path = "C:\\Users\\admin\Desktop\\YT\\Resume Scanner\\test"  # Folder containing the resumes
    output_file = "test.xlsx"  # Output Excel file

    resume_data = process_resumes(folder_path)
    create_excel(resume_data, output_file)

    print(f"Data from resumes has been successfully saved to {output_file}")

if __name__ == "__main__":
    main()